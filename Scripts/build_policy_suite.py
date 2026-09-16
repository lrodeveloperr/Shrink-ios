#!/usr/bin/env python3
"""Create the bilingual policy suite from the retained policy template."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]


def shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    properties.append(fill)


def configure(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    section = doc.sections[0]
    doc.settings.odd_and_even_pages_header_footer = False
    section.different_first_page_header_footer = True
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    section.header.paragraphs[0].clear()
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(38, 41, 47)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0
    for name, size, color in (("Title", 28, "12305C"), ("Heading 1", 21, "12305C"), ("Heading 2", 13, "1769AA")):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("WorksBien Studios Inc. · Shrinkflation Price Scanner · September 16, 2026")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 105, 115)


def add_markdown(doc: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line.replace("`", ""))


def add_status_table(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Release integration record", level=1)
    doc.add_paragraph(
        "This appendix records how the policy modules map to the reviewed app and release materials. "
        "It is an operational record, not a substitute for legal advice."
    )
    rows = [
        ("Market", "United States, including Puerto Rico; English and Puerto Rico Spanish only", "Complete"),
        ("Local shopping data", "Stored in the app container; no WorksBien account or runtime product API", "Complete"),
        ("Catalogue", "USDA FoodData Central branded-food release dated April 30, 2026; 426,044 selected products", "Complete"),
        ("Advertising", "Google Mobile Ads, UMP/ATT when applicable, and a WorksBien fallback banner", "Production configuration pending"),
        ("Purchase", "StoreKit 2 non-consumable Remove Banner entitlement and restore flow", "Complete in code; sandbox test pending"),
        ("Public routes", "English and Puerto Rico Spanish privacy and terms routes", "Prepared for deployment"),
        ("TestFlight", "Bilingual beta description, test focus, and review notes", "Prepared; intentionally not wired"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    headers = ("Module", "Code-grounded position", "Release state")
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        shade(cell, "12305C")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for module, position, status in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, (module, position, status)):
            cell.text = value
    doc.add_heading("Source anchors", level=2)
    for item in (
        "Apple App Review Guidelines and Standard Licensed Application End User License Agreement",
        "Apple App Store Connect app-information and App Privacy guidance",
        "Google Mobile Ads SDK data-disclosure and User Messaging Platform guidance",
        "USDA FoodData Central branded-food data source",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Human release gates", level=2)
    for item in (
        "Replace Google's sample advertising identifiers only in the approved production archive.",
        "Reconcile the final SDK privacy manifests, App Privacy answers, ATT behavior, and UMP configuration.",
        "Supply current App Review and TestFlight feedback contacts.",
        "Complete device, StoreKit sandbox, accessibility, media, signing, and archive checks.",
        "Upload or submit TestFlight only after a separate deliberate authorization.",
    ):
        doc.add_paragraph(item, style="List Bullet")


def build(template: Path, output: Path) -> None:
    doc = Document(template)
    configure(doc)
    add_footer(doc.sections[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Shrinkflation Price Scanner\nPolicy Suite")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "US English · Puerto Rico Spanish\n"
        "Version 1.0 (8) · Effective September 16, 2026\n"
        "Global App Policy Lego System"
    ).bold = True
    doc.add_paragraph(
        "Code-grounded privacy and terms modules for the United States, including Puerto Rico. "
        "The public web pages and in-app links use these same controlled sources."
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    for index, (label, path) in enumerate((
        ("US English", ROOT / "Policies" / "en-US" / "privacy.md"),
        ("US English", ROOT / "Policies" / "en-US" / "terms.md"),
        ("Español de Puerto Rico", ROOT / "Policies" / "es-PR" / "privacy.md"),
        ("Español de Puerto Rico", ROOT / "Policies" / "es-PR" / "terms.md"),
    )):
        doc.add_page_break()
        eyebrow = doc.add_paragraph()
        eyebrow.add_run(label.upper()).bold = True
        add_markdown(doc, path)

    add_status_table(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.template, args.output)


if __name__ == "__main__":
    main()
